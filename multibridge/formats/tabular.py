"""CSV and Excel: one row per data row of the main table. Other tables are noted, not joined.

As a destination, the example's main table gives the columns (in order), their types and how
values are written (date style, decimal comma, the CSV separator). Custom fields carried through
become extra columns at the end.
"""

from __future__ import annotations

import csv
import datetime as dt
import io

from ..output import EXTRAS, extra_names, render
from ..parsers import csv_, excel
from . import Destination, Format

BOM = b"\xef\xbb\xbf"
EXCEL_DATES = {"%d/%m/%Y": "DD/MM/YYYY", "%d-%m-%Y": "DD-MM-YYYY", "%d.%m.%Y": "DD.MM.YYYY", "%Y%m%d": "YYYYMMDD"}


def _columns(fmt: Format, data: bytes, name: str, pack) -> tuple[str, list]:
    from ..expected import target_from_column
    from ..profile import profile_dataset

    ds = profile_dataset(fmt.read(data, name), packs=None)
    root = fmt.root(ds)
    if root is None or not ds.tables[root]:
        raise ValueError(f"{name}: el ejemplo necesita una cabecera y al menos una fila de ejemplo")
    targets = [target_from_column(f.path, f, segments=(f.path,)) for f in ds.fields if f.entity == root]
    return root, targets


def _rows(items: list[dict], expected) -> tuple[list[str], list[list]]:
    extras = extra_names(items)
    header = [t.path for t in expected.targets] + extras
    rows = [[item.get(t.path) for t in expected.targets] + [(item.get(EXTRAS) or {}).get(e) for e in extras]
            for item in items]
    return header, rows


class Csv(Format):
    name, label, extensions, mime = "csv", "CSV", (".csv", ".tsv", ".txt"), "text/csv"

    def read(self, data, name):
        return csv_.parse(data, name)

    def destination(self, data, name, pack, options):
        _, targets = _columns(self, data, name, pack)
        text, _ = csv_.decode(data)
        return Destination(self.name, targets, template={"delimiter": csv_.sniff_delimiter(text),
                                                         "bom": data.startswith(BOM)})

    def write(self, items, expected):
        tpl = expected.template or {}
        header, rows = _rows(items, expected)
        by_path = {t.path: t for t in expected.targets}
        out = io.StringIO()
        w = csv.writer(out, delimiter=tpl.get("delimiter", ";"), lineterminator="\n")
        w.writerow(header)
        for row in rows:
            w.writerow([render(by_path[h], v) if h in by_path else ("" if v is None else str(v))
                        for h, v in zip(header, row)])
        return (BOM if tpl.get("bom") else b"") + out.getvalue().encode("utf-8")

    def check(self, data):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            return ["no es texto UTF-8"]
        lines = [l for l in csv.reader(io.StringIO(text), delimiter=csv_.sniff_delimiter(text)) if l]
        if not lines:
            return ["el CSV está vacío"]
        width = len(lines[0])
        bad = [n for n, l in enumerate(lines[1:], start=2) if len(l) != width]
        return [f"filas con un número de columnas distinto de la cabecera: {', '.join(map(str, bad[:10]))}"] if bad else []


class Excel(Format):
    name, label, extensions = "excel", "Excel", (".xlsx", ".xlsm")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def sniff(self, data, name):
        return super().sniff(data, name) or data[:4] == b"PK\x03\x04"

    def read(self, data, name):
        return excel.parse(data, name)

    def destination(self, data, name, pack, options):
        root, targets = _columns(self, data, name, pack)
        return Destination(self.name, targets, template={"sheet": root.rsplit(" (tabla", 1)[0][:31]})

    def write(self, items, expected):
        from openpyxl import Workbook

        header, rows = _rows(items, expected)
        by_path = {t.path: t for t in expected.targets}
        wb = Workbook()
        ws = wb.active
        ws.title = (expected.template or {}).get("sheet") or "Datos"
        ws.append(header)
        for r, row in enumerate(rows, start=2):
            for c, (h, v) in enumerate(zip(header, row), start=1):
                t = by_path.get(h)
                cell = ws.cell(row=r, column=c)
                if v is None:
                    continue
                if t is not None and t.type == "date":
                    try:
                        cell.value = dt.date.fromisoformat(str(v))
                        cell.number_format = EXCEL_DATES.get((t.render or {}).get("date"), "YYYY-MM-DD")
                        continue
                    except ValueError:
                        pass
                cell.value = v if isinstance(v, (int, float, str, bool)) else (
                    render(t, v) if t is not None else str(v))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def check(self, data):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True)
        except Exception as e:  # openpyxl raises many kinds for a broken file
            return [f"no es un Excel válido: {type(e).__name__}"]
        ws = wb.worksheets[0]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        return [] if header and any(h is not None for h in header) else ["la primera hoja no tiene cabecera"]


CSV, EXCEL = Csv(), Excel()
