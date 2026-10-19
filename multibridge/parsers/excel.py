"""Excel (.xlsx/.xlsm) as people actually make them.

Handles: several sheets, title rows above the header, merged group headers,
several tables per sheet, footers, cell comments used as field documentation,
formulas without a cached value, hidden sheets/columns, real Excel dates.
"""

from __future__ import annotations

import datetime as dt
import re
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter

from ..model import Dataset
from .table import add_tables

COMMENT_AUTHOR = re.compile(r"^[^:\n]{1,40}:\n")


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(timespec="minutes")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 10))
    return str(v).strip()


def parse(data: bytes, name: str) -> Dataset:
    values_wb = openpyxl.load_workbook(BytesIO(data), data_only=True)   # cached results
    formula_wb = openpyxl.load_workbook(BytesIO(data), data_only=False)  # formulas + comments
    ds = Dataset(source=name, format="excel")

    for ws in values_wb.worksheets:
        wf = formula_wb[ws.title]
        if ws.sheet_state != "visible":
            ds.notes.append(f"Hoja '{ws.title}' está oculta; se ha leído igualmente.")

        grid = [[cell_text(c.value) for c in row] for row in ws.iter_rows()]
        if not any(any(r) for r in grid):
            continue

        # Merged cells only store a value in the top-left cell: spread it.
        for rng in ws.merged_cells.ranges:
            top = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row - 1, rng.max_row):
                for c in range(rng.min_col - 1, rng.max_col):
                    if r < len(grid) and c < len(grid[r]):
                        grid[r][c] = top

        descriptions: dict[int, list[str]] = {}
        uncached: dict[int, int] = {}
        for row in wf.iter_rows():
            for cell in row:
                col = cell.column - 1
                if cell.comment:
                    text = COMMENT_AUTHOR.sub("", cell.comment.text).strip()
                    descriptions.setdefault(col, []).append(text)
                if cell.data_type == "f":
                    if col not in uncached:
                        descriptions.setdefault(col, []).append(f"Fórmula: {cell.value}")
                        uncached[col] = 0
                    if ws.cell(cell.row, cell.column).value is None:
                        uncached[col] = uncached.get(col, 0) + 1

        for col, n in uncached.items():
            if n:
                ds.notes.append(
                    f"Hoja '{ws.title}', columna {get_column_letter(col + 1)}: {n} fórmula(s) sin valor "
                    "calculado (el archivo no se guardó desde Excel). Ábrelo y guárdalo en Excel para obtener los valores."
                )
        hidden = [k for k, d in ws.column_dimensions.items() if d.hidden]
        if hidden:
            ds.notes.append(f"Hoja '{ws.title}': columnas ocultas {', '.join(hidden)}; se han leído igualmente.")

        add_tables(ds, ws.title, grid, {c: " · ".join(d) for c, d in descriptions.items()})
    return ds
