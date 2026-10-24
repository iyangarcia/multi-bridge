"""Harder synthetic test files, all rendered from the same three invented patients.

    python examples/hard/build_hard.py

Because the patients are known, every file comes with its correct answer (answers.json):
    plan      the right decision per destination field: columns + transforms, "no_data",
              or "gap" (a correct recipe can't be expressed yet: a missing feature)
    records   exactly what the pipeline must output with that plan
    rejected  rows that must be set aside, and why

h1  English CSV: commas + BOM, quoted commas, Female/Male, grams, metres, spaced phones,
    a postal code that lost its leading zero, "A Rh+" / "O negativo", an impossible date
h2  Cryptic Spanish CSV: Windows-1252, a title row, NHC;APE1;APE2;NOM;FNAC;SX..., two surname
    columns, V/M for sex, dd.mm.yyyy, height "1,68" in metres, all in capitals
h3  HL7 v2.5: MLLP framing, mixed line breaks, DNI and NHC as repetitions of PID-3 (told apart
    only by PID-3.5 NI/MR), alias names, second given name, v2.5 phone components, hex escapes,
    a different Z-segment, missing segments
h4  FHIR: official + nickname names, home + mobile phones, a partial birth date, the doctor in
    generalPractitioner, weight/height as Observations, identifiers by system
h5  Excel: data on the second sheet, mixed date formats and numbers-as-text in one column,
    a combined "Nombre completo" column, a hidden DNI column, footers
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
ANSWERS: dict[str, dict] = {}


def record(**kw) -> dict:
    """An output record in the shape of expected/cdss_paciente.example.json."""
    return {
        "paciente": {k: kw.get(k) for k in ("id", "dni", "nombre", "apellidos", "fecha_nacimiento", "sexo", "telefono")},
        "direccion": {k: kw.get(k) for k in ("calle", "ciudad", "codigo_postal")},
        "datos_clinicos": {k: kw.get(k) for k in ("grupo_sanguineo", "peso_kg", "altura_cm", "alergia",
                                                  "medico_responsable")},
    }


def m(*cols, t=None) -> dict:
    return {"status": "mapped", "from": list(cols), "transforms": t or []}


NO = {"status": "no_data", "from": [], "transforms": []}


def gap(why: str) -> dict:
    return {"status": "gap", "why": why}


# --- h1: English CSV ------------------------------------------------------------------------------

def h1() -> None:
    header = ["Patient ID", "First name", "Last name", "DOB", "Gender", "Mobile", "Address", "City", "Zip",
              "Blood type", "Weight (g)", "Height (m)", "Allergies", "GP"]
    rows = [
        ["P-001", "Ana", "Pérez Gil", "1985-04-12", "Female", "600 11 22 33", '"Calle Uría, 12"', "Alicante", "3001",
         "A Rh+", "65500", "1.68", '"Penicilina; Látex"', "Dr. Luis Prendes"],
        ["P-002", "Javier", "Muñoz Fernández", "1960-11-05", "Male", "600 44 55 66", "Av. Galicia 4", "Gijón", "33212",
         "O negativo", "72500", "1.72", "", "Dra. Carmen Fueyo"],
        ["P-003", "Lucía María", "Álvarez Castro", "2001-09-07", "Female", "600 77 88 99", "C/ Mayor 1", "Avilés",
         "33402", "AB Rh+", "54000", "1.60", "Marisco", "Dra. Ana Rubiera"],
        ["P-004", "Test", "Fecha Imposible", "1990-02-30", "Unknown", "600 00 00 00", "C/ Falsa 1", "Oviedo", "33001",
         "B+", "80000", "1.80", "", "Dr. Luis Prendes"],
    ]
    text = ",".join(header) + "\r\n" + "\r\n".join(",".join(r) for r in rows) + "\r\n\r\n"
    (HERE / "h1_ingles_comas.csv").write_bytes(b"\xef\xbb\xbf" + text.encode("utf-8"))
    ANSWERS["h1_ingles_comas.csv"] = {
        "plan": {
            "paciente.id": m("Patient ID"), "paciente.dni": NO, "paciente.nombre": m("First name"),
            "paciente.apellidos": m("Last name"), "paciente.fecha_nacimiento": m("DOB"),
            "paciente.sexo": m("Gender", t=[{"op": "translate", "values": {"Female": "female", "Male": "male",
                                                                           "Unknown": "unknown"}}]),
            "paciente.telefono": m("Mobile", t=[{"op": "phone", "country": "34"}]),
            "direccion.calle": m("Address"), "direccion.ciudad": m("City"),
            "direccion.codigo_postal": m("Zip", t=[{"op": "pad", "width": 5}]),
            "datos_clinicos.grupo_sanguineo": m("Blood type", t=[{"op": "translate", "values": {
                "A Rh+": "A+", "O negativo": "O-", "AB Rh+": "AB+"}}]),
            "datos_clinicos.peso_kg": m("Weight (g)", t=[{"op": "number"}, {"op": "scale", "factor": 0.001}]),
            "datos_clinicos.altura_cm": m("Height (m)", t=[{"op": "number"}, {"op": "scale", "factor": 100}]),
            "datos_clinicos.alergia": m("Allergies"), "datos_clinicos.medico_responsable": m("GP"),
        },
        "records": [
            record(id="P-001", nombre="Ana", apellidos="Pérez Gil", fecha_nacimiento="1985-04-12", sexo="female",
                   telefono="+34600112233", calle="Calle Uría, 12", ciudad="Alicante", codigo_postal="03001",
                   grupo_sanguineo="A+", peso_kg=65.5, altura_cm=168, alergia="Penicilina; Látex",
                   medico_responsable="Dr. Luis Prendes"),
            record(id="P-002", nombre="Javier", apellidos="Muñoz Fernández", fecha_nacimiento="1960-11-05", sexo="male",
                   telefono="+34600445566", calle="Av. Galicia 4", ciudad="Gijón", codigo_postal="33212",
                   grupo_sanguineo="O-", peso_kg=72.5, altura_cm=172, medico_responsable="Dra. Carmen Fueyo"),
            record(id="P-003", nombre="Lucía María", apellidos="Álvarez Castro", fecha_nacimiento="2001-09-07",
                   sexo="female", telefono="+34600778899", calle="C/ Mayor 1", ciudad="Avilés", codigo_postal="33402",
                   grupo_sanguineo="AB+", peso_kg=54, altura_cm=160, alergia="Marisco",
                   medico_responsable="Dra. Ana Rubiera"),
        ],
        "rejected": [{"row": 3, "field": "paciente.fecha_nacimiento"}],
    }


# --- h2: cryptic Spanish CSV --------------------------------------------------------------------

def h2() -> None:
    lines = [
        "LISTADO PACIENTES - URGENCIAS;;;;;;;;;;;;;;",
        ";;;;;;;;;;;;;;",
        "NHC;APE1;APE2;NOM;FNAC;SX;TLF;DOM;LOC;CP;GS;PESO;TALLA;ALERG;MED",
        "10001;PÉREZ;GIL;ANA;12.04.1985;M;600112233;CALLE URÍA, 12;ALICANTE;03001;A+;65,5;1,68;PENICILINA;LUIS PRENDES",
        "10002;MUÑOZ;FERNÁNDEZ;JAVIER;05.11.1960;V;600445566;AV. GALICIA 4;GIJÓN;33212;0-;72,5;1,72;;CARMEN FUEYO",
        "10003;ÁLVAREZ;CASTRO;LUCÍA MARÍA;07.09.2001;M;600778899;C/ MAYOR 1;AVILÉS;33402;AB+;54;1,60;MARISCO;ANA RUBIERA",
    ]
    (HERE / "h2_cifrado.csv").write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1252"))
    title = [{"op": "title"}]
    ANSWERS["h2_cifrado.csv"] = {
        "plan": {
            "paciente.id": m("NHC"), "paciente.dni": NO, "paciente.nombre": m("NOM", t=title),
            "paciente.apellidos": m("APE1", "APE2", t=[{"op": "join_columns", "sep": " "}, {"op": "title"}]),
            "paciente.fecha_nacimiento": m("FNAC", t=[{"op": "date", "from": "%d.%m.%Y"}]),
            "paciente.sexo": m("SX", t=[{"op": "translate", "values": {"M": "female", "V": "male"}}]),
            "paciente.telefono": m("TLF", t=[{"op": "phone", "country": "34"}]),
            "direccion.calle": m("DOM"), "direccion.ciudad": m("LOC"), "direccion.codigo_postal": m("CP"),
            "datos_clinicos.grupo_sanguineo": m("GS", t=[{"op": "translate", "values": {"0-": "O-"}}]),
            "datos_clinicos.peso_kg": m("PESO", t=[{"op": "number", "decimal": ","}]),
            "datos_clinicos.altura_cm": m("TALLA", t=[{"op": "number", "decimal": ","}, {"op": "scale", "factor": 100}]),
            "datos_clinicos.alergia": m("ALERG"), "datos_clinicos.medico_responsable": m("MED", t=title),
        },
        "records": [
            record(id="10001", nombre="Ana", apellidos="Pérez Gil", fecha_nacimiento="1985-04-12", sexo="female",
                   telefono="+34600112233", calle="CALLE URÍA, 12", ciudad="ALICANTE", codigo_postal="03001",
                   grupo_sanguineo="A+", peso_kg=65.5, altura_cm=168, alergia="PENICILINA",
                   medico_responsable="Luis Prendes"),
            record(id="10002", nombre="Javier", apellidos="Muñoz Fernández", fecha_nacimiento="1960-11-05", sexo="male",
                   telefono="+34600445566", calle="AV. GALICIA 4", ciudad="GIJÓN", codigo_postal="33212",
                   grupo_sanguineo="O-", peso_kg=72.5, altura_cm=172, medico_responsable="Carmen Fueyo"),
            record(id="10003", nombre="Lucía María", apellidos="Álvarez Castro", fecha_nacimiento="2001-09-07",
                   sexo="female", telefono="+34600778899", calle="C/ MAYOR 1", ciudad="AVILÉS", codigo_postal="33402",
                   grupo_sanguineo="AB+", peso_kg=54, altura_cm=160, alergia="MARISCO",
                   medico_responsable="Ana Rubiera"),
        ],
        "rejected": [],
    }


# --- h3: HL7 v2.5 -------------------------------------------------------------------------------

def h3() -> None:
    m1 = [
        "MSH|^~\\&|HIS|HOSP_ALC|PRAGMATECH|CDSS|20260401090000||ADT^A04^ADT_A01|H3-001|P|2.5",
        "PID|1||12345678Z^^^MIR^NI~NHC-30001^^^HALC^MR||PÉREZ GIL^ANA^^^^^L~PEREZ^ANITA^^^^^N||198504121030|F|||"
        "Calle Uría, 12^^Alicante^Alicante^03001^ESP~Av. Trabajo 1^^Alicante^^03002^ESP^B||"
        "^PRN^PH^^34^^600112233~^NET^Internet^ana@example.com",
        "PV1|1|O|CEX^101^1||||D0100^PRENDES^LUIS",
        "AL1|1|DA|PEN^Penicilina|SV",
        "AL1|2|MA|LAT^Látex|MO",
        "ZBG|1|A POS|65.5|168",
    ]
    m2 = [
        "MSH|^~\\&|HIS|HOSP_ALC|PRAGMATECH|CDSS|20260401091500||ADT^A04^ADT_A01|H3-002|P|2.5",
        "PID|1||NHC-30002^^^HALC^MR||MU\\XD1\\OZ FERN\\XC1\\NDEZ^JAVIER||19601105|M|||Av. Galicia 4^^Gijón^Asturias^33212^ESP||"
        "^PRN^PH^^34^^600445566",
        "PV1|1|I|MIR^202^2||||D0200^FUEYO^CARMEN",
        "ZBG|1|0 NEG|72.5|172",
    ]
    m3 = [
        "MSH|^~\\&|HIS|HOSP_ALC|PRAGMATECH|CDSS|20260401093000||ADT^A04^ADT_A01|H3-003|P|2.5",
        "PID|1||NHC-30003^^^HALC^MR~53123456T^^^MIR^NI||ÁLVAREZ CASTRO^LUCÍA^MARÍA||20010907|F|||"
        "C/ Mayor 1^^Avilés^Asturias^33402^ESP||^PRN^PH^^34^^600778899",
        "AL1|1|FA|MAR^Marisco|MO",
    ]
    data = b"\x0b" + "\r".join(m1).encode("utf-8") + b"\x1c\r"           # MLLP framing, \r separators
    data += b"\x0b" + "\r\n".join(m2).encode("utf-8") + b"\x1c\r"        # \r\n separators
    data += "\n".join(m3).encode("utf-8") + b"\n"                         # no framing, \n separators
    (HERE / "h3_hl7_v25.hl7").write_bytes(data)
    ANSWERS["h3_hl7_v25.hl7"] = {
        "plan": {
            # PID-3 holds DNI and NHC as repetitions; PID-3.5 (identifier type) tells them apart
            "paciente.id": m("PID-3.1 [PID-3.5=MR]"), "paciente.dni": m("PID-3.1 [PID-3.5=NI]"),
            "paciente.nombre": m("PID-5.2", "PID-5.3", t=[{"op": "join_columns", "sep": " "}, {"op": "title"}]),
            "paciente.apellidos": m("PID-5.1", t=[{"op": "title"}, {"op": "first"}]),
            "paciente.fecha_nacimiento": m("PID-7", t=[{"op": "date", "from": "%Y%m%d"}]),
            "paciente.sexo": m("PID-8", t=[{"op": "translate", "values": {"F": "female", "M": "male"}}]),
            "paciente.telefono": m("PID-13.7", t=[{"op": "phone", "country": "34"}]),
            "direccion.calle": m("PID-11.1", t=[{"op": "first"}]),
            "direccion.ciudad": m("PID-11.3", t=[{"op": "first"}]),
            "direccion.codigo_postal": m("PID-11.5", t=[{"op": "first"}]),
            "datos_clinicos.grupo_sanguineo": m("ZBG-2", t=[{"op": "translate", "values": {"A POS": "A+", "0 NEG": "O-"}}]),
            "datos_clinicos.peso_kg": m("ZBG-3", t=[{"op": "number"}]),
            "datos_clinicos.altura_cm": m("ZBG-4", t=[{"op": "number"}]),
            "datos_clinicos.alergia": m("AL1-3.2", t=[{"op": "join_values", "sep": "; "}]),
            "datos_clinicos.medico_responsable": m("PV1-7.3", "PV1-7.2", t=[{"op": "join_columns", "sep": " "},
                                                                          {"op": "title"}]),
        },
        "records": [
            record(id="NHC-30001", dni="12345678Z", nombre="Ana", apellidos="Pérez Gil",
                   fecha_nacimiento="1985-04-12", sexo="female", telefono="+34600112233", calle="Calle Uría, 12", ciudad="Alicante", codigo_postal="03001", grupo_sanguineo="A+",
                   peso_kg=65.5, altura_cm=168, alergia="Penicilina; Látex", medico_responsable="Luis Prendes"),
            record(id="NHC-30002", nombre="Javier", apellidos="Muñoz Fernández", fecha_nacimiento="1960-11-05",
                   sexo="male",
                   telefono="+34600445566", calle="Av. Galicia 4", ciudad="Gijón", codigo_postal="33212",
                   grupo_sanguineo="O-", peso_kg=72.5, altura_cm=172, medico_responsable="Carmen Fueyo"),
            record(id="NHC-30003", dni="53123456T", nombre="Lucía María", apellidos="Álvarez Castro",
                   fecha_nacimiento="2001-09-07", sexo="female",
                   telefono="+34600778899",
                   calle="C/ Mayor 1", ciudad="Avilés", codigo_postal="33402", alergia="Marisco"),
        ],
        "rejected": [],
    }


# --- h4: FHIR -----------------------------------------------------------------------------------

def h4() -> None:
    def patient(pid, nhc, dni, official, nick, birth, gender, home, mobile, line, city, cp, gp):
        return {"resourceType": "Patient", "id": pid,
                "identifier": [{"system": "urn:hospital:nhc", "type": {"coding": [{"code": "MR"}]}, "value": nhc},
                               {"system": "urn:es:dni", "type": {"coding": [{"code": "NI"}]}, "value": dni}],
                "name": [{"use": "official", "family": official[0], "given": official[1]}]
                        + ([{"use": "nickname", "given": [nick]}] if nick else []),
                "telecom": [{"system": "phone", "use": "home", "value": home},
                            {"system": "phone", "use": "mobile", "value": mobile},
                            {"system": "email", "value": f"{pid}@example.com"}],
                "gender": gender, "birthDate": birth,
                "address": [{"use": "home", "line": [line], "city": city, "postalCode": cp, "country": "ES"}],
                "generalPractitioner": [{"reference": f"Practitioner/{gp[0]}", "display": gp[1]}]}

    def obs(oid, pid, code, value, unit):
        return {"resourceType": "Observation", "id": oid, "status": "final",
                "code": {"coding": [{"system": "http://loinc.org", "code": code}]},
                "subject": {"reference": f"Patient/{pid}"}, "valueQuantity": {"value": value, "unit": unit}}

    resources = [
        patient("a", "NHC-40001", "12345678Z", ("Pérez Gil", ["Ana"]), "Anita", "1985-04-12", "female",
                "+34985111222", "+34600112233", "Calle Uría, 12", "Alicante", "03001", ("d1", "Dr. Luis Prendes")),
        patient("b", "NHC-40002", "87654321X", ("Muñoz Fernández", ["Javier"]), None, "1960-11-05", "male",
                "+34985333444", "+34600445566", "Av. Galicia 4", "Gijón", "33212", ("d2", "Dra. Carmen Fueyo")),
        patient("c", "NHC-40003", "53123456T", ("Álvarez Castro", ["Lucía", "María"]), "Lu", "2001-09", "female",
                "+34985555666", "+34600778899", "C/ Mayor 1", "Avilés", "33402", ("d3", "Dra. Ana Rubiera")),
        obs("o1", "a", "29463-7", 65.5, "kg"), obs("o2", "a", "8302-2", 168, "cm"),
        obs("o3", "b", "29463-7", 72.5, "kg"), obs("o4", "b", "8302-2", 172, "cm"),
        {"resourceType": "AllergyIntolerance", "id": "al1", "patient": {"reference": "Patient/a"},
         "code": {"coding": [{"system": "http://snomed.info/sct", "code": "764146007"}], "text": "Penicilina"}},
    ]
    bundle = {"resourceType": "Bundle", "type": "collection", "entry": [{"resource": r} for r in resources]}
    (HERE / "h4_fhir_completo.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    ANSWERS["h4_fhir_completo.json"] = {
        "plan": {
            "paciente.id": m("identifier[nhc].value"), "paciente.dni": m("identifier[dni].value"),
            "paciente.nombre": m("name[official].given", t=[{"op": "join_values", "sep": " "}]),
            "paciente.apellidos": m("name[official].family"), "paciente.fecha_nacimiento": m("birthDate"),
            "paciente.sexo": m("gender"),
            "paciente.telefono": m("telecom[phone-mobile].value"),   # a mobile reaches the patient; home is the rival
            "direccion.calle": m("address.line"), "direccion.ciudad": m("address.city"),
            "direccion.codigo_postal": m("address.postalCode"), "datos_clinicos.grupo_sanguineo": NO,
            "datos_clinicos.peso_kg": m("Observation.valueQuantity.value [Observation.code.coding.code=29463-7]",
                                        t=[{"op": "number"}]),
            "datos_clinicos.altura_cm": m("Observation.valueQuantity.value [Observation.code.coding.code=8302-2]",
                                          t=[{"op": "number"}]),
            "datos_clinicos.alergia": m("AllergyIntolerance.code.text"),
            "datos_clinicos.medico_responsable": m("generalPractitioner.display"),
        },
        "records": [
            record(id="NHC-40001", dni="12345678Z", nombre="Ana", telefono="+34600112233", apellidos="Pérez Gil", fecha_nacimiento="1985-04-12", sexo="female",
                   calle="Calle Uría, 12", ciudad="Alicante", codigo_postal="03001", alergia="Penicilina",
                   peso_kg=65.5, altura_cm=168,
                   medico_responsable="Dr. Luis Prendes"),
            record(id="NHC-40002", dni="87654321X", nombre="Javier", telefono="+34600445566",
                   apellidos="Muñoz Fernández", fecha_nacimiento="1960-11-05",
                   sexo="male", calle="Av. Galicia 4", ciudad="Gijón", codigo_postal="33212",
                   peso_kg=72.5, altura_cm=172, medico_responsable="Dra. Carmen Fueyo"),
        ],
        "rejected": [{"row": 2, "field": "paciente.fecha_nacimiento"}],   # "2001-09": FHIR allows it, ISO date doesn't
    }


# --- h5: Excel ----------------------------------------------------------------------------------

def h5() -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary.append(["Servicio", "Pacientes"])
    summary.append(["Urgencias", 3])
    data = wb.create_sheet("Pacientes")
    data["A1"] = "Exportación de pacientes (datos sintéticos)"
    data.merge_cells("A1:L1")
    data.append([])
    data.append(["Identificación", None, None, None, "Clínica", None, None, None, None, None, None, None])
    data.merge_cells("A3:D3")
    data.merge_cells("E3:L3")
    data.append(["NHC", "Nombre completo", "DNI", "Fecha nacimiento", "Sexo", "Peso (kg)", "Altura (cm)",
                 "Grupo sanguíneo", "Alergias", "Teléfono", "Ciudad", "CP"])
    data.append(["E-1", "Ana Pérez Gil", "12345678Z", dt.date(1985, 4, 12), "Mujer", 65.5, 168, "A+", "Penicilina",
                 "600112233", "Alicante", "03001"])
    data.append(["E-2", "Javier Muñoz Fernández", "87654321X", "05/11/1960", "Hombre", "72,5", "172", "0-", None,
                 "600445566", "Gijón", "33212"])
    data.append(["E-3", "Lucía María Álvarez Castro", "53123456T", "2001-09-07", "Mujer", 54, 160, "AB+", "Marisco",
                 "600778899", "Avilés", "33402"])
    data.append([])
    data.append(["Total: 3"])
    data.column_dimensions["C"].hidden = True
    for row in range(5, 8):
        data.cell(row, 4).number_format = "DD/MM/YYYY"
    wb.save(HERE / "h5_excel_dificil.xlsx")
    ANSWERS["h5_excel_dificil.xlsx"] = {
        "plan": {
            "paciente.id": m("Identificación / NHC"), "paciente.dni": m("Identificación / DNI"),
            # 'Nombre completo' is split: given name(s), then two surnames (phase 7)
            "paciente.nombre": m("Identificación / Nombre completo",
                                 t=[{"op": "name_part", "part": "given", "surnames": 2}]),
            "paciente.apellidos": m("Identificación / Nombre completo",
                                    t=[{"op": "name_part", "part": "family", "surnames": 2}]),
            "paciente.fecha_nacimiento": m("Identificación / Fecha nacimiento",
                                           t=[{"op": "date", "from": ["%d/%m/%Y", "%Y-%m-%d"]}]),
            "paciente.sexo": m("Clínica / Sexo", t=[{"op": "translate", "values": {"Hombre": "male", "Mujer": "female"}}]),
            "paciente.telefono": m("Clínica / Teléfono", t=[{"op": "phone", "country": "34"}]),
            "direccion.calle": NO, "direccion.ciudad": m("Clínica / Ciudad"), "direccion.codigo_postal": m("Clínica / CP"),
            "datos_clinicos.grupo_sanguineo": m("Clínica / Grupo sanguíneo",
                                                t=[{"op": "translate", "values": {"0-": "O-"}}]),
            "datos_clinicos.peso_kg": m("Clínica / Peso (kg)", t=[{"op": "number", "decimal": ","}]),
            "datos_clinicos.altura_cm": m("Clínica / Altura (cm)", t=[{"op": "number"}]),
            "datos_clinicos.alergia": m("Clínica / Alergias"), "datos_clinicos.medico_responsable": NO,
        },
        "records": [
            record(id="E-1", dni="12345678Z", nombre="Ana", apellidos="Pérez Gil", fecha_nacimiento="1985-04-12", sexo="female", telefono="+34600112233", ciudad="Alicante",
                   codigo_postal="03001", grupo_sanguineo="A+", peso_kg=65.5, altura_cm=168, alergia="Penicilina"),
            record(id="E-2", dni="87654321X", nombre="Javier", apellidos="Muñoz Fernández", fecha_nacimiento="1960-11-05", sexo="male", telefono="+34600445566", ciudad="Gijón",
                   codigo_postal="33212", grupo_sanguineo="O-", peso_kg=72.5, altura_cm=172),
            record(id="E-3", dni="53123456T", nombre="Lucía María", apellidos="Álvarez Castro", fecha_nacimiento="2001-09-07", sexo="female", telefono="+34600778899", ciudad="Avilés",
                   codigo_postal="33402", grupo_sanguineo="AB+", peso_kg=54, altura_cm=160, alergia="Marisco"),
        ],
        "rejected": [],
    }


if __name__ == "__main__":
    for build in (h1, h2, h3, h4, h5):
        build()
    (HERE / "answers.json").write_text(json.dumps(ANSWERS, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"built {len(ANSWERS)} files + answers.json")
