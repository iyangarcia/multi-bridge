"""Regenerate the "messy" synthetic examples. All people, codes and IDs are invented.

    python examples/build_examples.py

hospital_c_messy.xlsx       Excel as a real admin would make it: title rows, merged group
                            headers, header comments, formulas, free-text diagnoses, an
                            unnamed column, a second table on the same sheet, a hidden sheet.
hospital_a_diagnosticos.hl7 HL7 v2.3 with repeated DG1 (coded + free text), OBX with LOINC,
                            repetitions (~), escape sequences and a custom Z-segment.
hospital_b_condiciones.json FHIR Bundle with Condition resources: ICD-10, SNOMED CT,
                            text-only diagnoses, a custom extension, an AllergyIntolerance.
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos"
    ws["A1"] = "Hospital Universitario de Cabueñes — Listado de ingresos (datos sintéticos)"
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Exportado: 12/03/2026 por SIH-Cabueñes"

    for rng, label in [("A4:D4", "Paciente"), ("E4:G4", "Antropometría"), ("H4:J4", "Clínica")]:
        ws.merge_cells(rng)
        cell = ws[rng.split(":")[0]]
        cell.value = label
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    headers = ["NHC", "Nombre y apellidos", "F. Nac.", "Sexo", "Peso (kg)", "Talla (cm)", "IMC",
               "Diagnóstico principal", "Otros diagnósticos", "Campo libre 1"]
    for col, h in enumerate(headers, start=1):
        ws.cell(5, col, h).font = Font(bold=True)
    ws["H5"].comment = Comment("Admisión:\nTexto libre del médico; a veces escriben el código CIE-10", "Admisión")
    ws["J5"].comment = Comment("Usado por Urgencias: nivel de triaje Manchester (1 rojo … 5 azul)", "Urgencias")

    rows = [
        ("C-0101", "Marta Iglesias Prado", dt.date(1961, 3, 4), "M", 71.5, 160, "DM2 mal controlada", "HTA; DLP", 3, "R"),
        ("C-0102", "Luis Ordóñez Vega", dt.date(1948, 11, 20), "H", 84, 172, "Reagudización de EPOC", "AP de IAM", 2, None),
        ("C-0103", "Rosa Fanjul Díaz", dt.date(1979, 6, 1), "M", 58.2, 165, "Sospecha de NAC", None, 3, None),
        ("C-0104", "Iván Castañón Suárez", dt.date(1990, 1, 15), "H", 77, 180, "IRA", "DiabA", 2, "R"),
        ("C-0105", "Carmen Rubio Llano", dt.date(1955, 8, 30), "M", 66.4, 158, "J18.9",
         "insuficiencia cardiaca congestiva", 1, None),
        ("C-0106", "Tomás Valdés Menéndez", dt.date(1972, 12, 9), "H", 92.3, 175, "diabetis tipo 2", "obesidad", 4, None),
    ]
    for r, (nhc, name, born, sex, kg, cm, dx, other, triage, flag) in enumerate(rows, start=6):
        for col, v in enumerate([nhc, name, born, sex, kg, cm, None, dx, other, triage, flag], start=1):
            ws.cell(r, col, v)
        ws.cell(r, 3).number_format = "DD/MM/YYYY"
        ws.cell(r, 7, f"=E{r}/(F{r}/100)^2")  # never calculated: openpyxl does not evaluate formulas
    ws.column_dimensions["K"].hidden = True    # someone's private marker column, no header

    ws["A13"] = "Total pacientes: 6"
    ws["A14"] = "Nota: datos sintéticos generados para la demo"
    ws["A16"], ws["B16"] = "Código", "Significado"
    ws["A17"], ws["B17"] = "H", "Hombre"
    ws["A18"], ws["B18"] = "M", "Mujer"

    legend = wb.create_sheet("Leyenda triaje")
    legend.append(["Nivel", "Color", "Tiempo máximo (min)"])
    for row in [(1, "Rojo", 0), (2, "Naranja", 10), (3, "Amarillo", 60), (4, "Verde", 120), (5, "Azul", 240)]:
        legend.append(row)
    legend.sheet_state = "hidden"
    wb.save(path)


HL7_MESSAGES = [
    [
        "MSH|^~\\&|HIS_HUCA|HUCA_OVIEDO|PRAGMATECH|CDSS|20260315083000||ADT^A01|MSG1001|P|2.3",
        "EVN|A01|20260315083000",
        "PID|1||NHC-20001^^^HUCA~ASTU123456^^^SESPA||PRIETO LLANO^ROSA||19580214|F|||"
        "C/ Uria 30^^Oviedo^Asturias^33003^ES||+34600111999~+34985000111||||||10293847K",
        "PV1|1|I|NEU^305^1||||D0457^MARTINEZ^PABLO",
        "DG1|1||J18.9^Neumonía, organismo no especificado^I10||20260315|A",
        "DG1|2|||Sospecha de NAC \\T\\ derrame pleural|20260315|W",
        "DG1|3|||AP de IAM\\.br\\HTA|20260315|A",
        "OBX|1|NM|29463-7^Peso corporal^LN||66.5|kg|||||F",
        "OBX|2|NM|8302-2^Talla^LN||162|cm|||||F",
        "ZTR|1|3^Amarillo|20260315082500",
    ],
    [
        "MSH|^~\\&|HIS_HUCA|HUCA_OVIEDO|PRAGMATECH|CDSS|20260315101500||ADT^A01|MSG1002|P|2.3",
        "EVN|A01|20260315101500",
        "PID|1||NHC-20002^^^HUCA||FERNANDEZ-VALLINA^MANUEL||19670901|M|||"
        "Av. Galicia 2^^Gijon^Asturias^33212^ES||+34600222888||||||10928374M",
        "PV1|1|E|URG^02^1||||D0112^SUAREZ^ANA",
        "DG1|1||E11^Diabetes mellitus tipo 2^I10||20260315|F",
        "DG1|2|||IRA|20260315|W",
        "DG1|3|||DiabA|20260315|W",
        "ZTR|1|2^Naranja|20260315101000",
    ],
    [
        "MSH|^~\\&|HIS_HUCA|HUCA_OVIEDO|PRAGMATECH|CDSS|20260316004500||ADT^A01|MSG1003|P|2.3",
        "EVN|A01|20260316004500",
        "PID|1||NHC-20003^^^HUCA||GARCIA NOVAL^LUCIA||19991230|F|||"
        "C/ Mayor 7^^Aviles^Asturias^33402^ES||+34600333777||||||71829304Z",
        "PV1|1|I|INF^110^2||||D0389^ALONSO^JORGE",
        "DG1|1||A41.9^Sepsis, no especificada^I10||20260316|A",
        "DG1|2|||ITU por E. coli|20260316|A",
        "AL1|1|DA|PEN^Penicilina|SV|Urticaria~Disnea",
        "ZTR|1|1^Rojo|20260316003000",
    ],
]


def build_hl7(path: Path) -> None:
    path.write_bytes("\r".join(seg for msg in HL7_MESSAGES for seg in msg).encode("utf-8") + b"\r")


def build_fhir(path: Path) -> None:
    def condition(cid, code=None, text=None, verification="confirmed", extension=None):
        res = {
            "resourceType": "Condition", "id": cid,
            "clinicalStatus": {"coding": [{"system": "http://terminology.hl7.org/CodeSystem/condition-clinical",
                                           "code": "active"}]},
            "verificationStatus": {"coding": [{"system": "http://terminology.hl7.org/CodeSystem/condition-ver-status",
                                               "code": verification}]},
            "code": {k: v for k, v in (("coding", code), ("text", text)) if v},
            "subject": {"reference": "Patient/p1"},
            "onsetDateTime": "2026-03-14",
        }
        if extension:
            res["extension"] = extension
        return res

    bundle = {"resourceType": "Bundle", "type": "collection", "entry": [{"resource": r} for r in [
        {"resourceType": "Patient", "id": "p1", "identifier": [{"system": "urn:hule:mrn", "value": "MRN-56001"}],
         "name": [{"family": "Robles Abad", "given": ["Inés"]}], "gender": "female", "birthDate": "1970-05-02"},
        condition("c1", code=[{"system": "http://hl7.org/fhir/sid/icd-10", "code": "E11",
                               "display": "Type 2 diabetes mellitus"}], text="DM2"),
        condition("c2", text="Sospecha de neumonía", verification="provisional"),
        condition("c3", code=[{"system": "http://snomed.info/sct", "code": "38341003",
                               "display": "Hypertensive disorder, systemic arterial"}], text="HTA"),
        condition("c4", text="EPOC", extension=[{
            "url": "http://hule.sacyl.es/fhir/StructureDefinition/servicio-origen", "valueString": "Neumología"}]),
        {"resourceType": "AllergyIntolerance", "id": "a1", "code": {"text": "Penicilina"},
         "criticality": "high", "patient": {"reference": "Patient/p1"}},
    ]]}
    path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    build_excel(HERE / "hospital_c_messy.xlsx")
    build_hl7(HERE / "hospital_a_diagnosticos.hl7")
    build_fhir(HERE / "hospital_b_condiciones.json")
    print("examples regenerated")
