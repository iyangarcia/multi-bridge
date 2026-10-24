"""The web app in English (web/i18n/en.json): the static pages and what the API says are all covered.

The page translates in the browser (web/js/comun.js, traducirTexto); `translate` below does the same
matching so a new Spanish text without an English version fails here instead of showing up in Spanish.
"""

import io
import json
import re
import shutil
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
import pytest

from multibridge.api import ApiError, App, route

ROOT = Path(__file__).parent.parent
EN = json.loads((ROOT / "web" / "i18n" / "en.json").read_text(encoding="utf-8"))
SPANISH = re.compile(r"[a-záéíóúñ]{3}", re.I)
ATTRIBUTES = ("title", "placeholder", "aria-label", "alt", "label")


def normal(s: str) -> str:
    return " ".join(s.split())


def _rule(es: str):
    parts = re.split(r"(\{\d(?::n)?\})", es)
    rx = "".join(r"(\d+(?:[.,]\d+)?)" if re.fullmatch(r"\{\d:n\}", p) else "(.+?)" if re.fullmatch(r"\{\d\}", p)
                 else re.escape(p) for p in parts)
    return re.compile(f"^{rx}$"), re.findall(r"\{(\d)(?::n)?\}", es)


RULES = [(*_rule(es), en) for es, en in EN["patrones"]]


def translate(text: str, level: int = 0) -> str:
    t = normal(text)
    if t in EN["textos"]:
        return EN["textos"][t]
    if level > 3:
        return t
    for rx, order, en in RULES:
        m = rx.match(t)
        if m:
            parts = {n: translate(m.group(i + 1), level + 1) for i, n in enumerate(order)}
            return re.sub(r"\{(\d)\}", lambda x: parts[x.group(1)], en)
    for sep in (" · ", "; ", ", ", " + "):
        if sep in t:
            pieces = [translate(x, level + 1) for x in t.split(sep)]
            if sep.join(pieces) != t:
                return sep.join(pieces)
    return t


def test_dictionary_is_well_formed():
    assert all(isinstance(v, str) and v.strip() for v in EN["textos"].values())
    assert all(k == normal(k) for k in list(EN["textos"]) + list(EN["html"]))
    for es, en in EN["patrones"]:
        assert set(re.findall(r"\{(\d)", en)) == set(re.findall(r"\{(\d)", es)), es


class _Page(HTMLParser):
    """Visible Spanish on a page: text, translatable attributes and whole data-i18n-html blocks."""

    SKIP = {"code", "pre", "script", "style", "textarea"}

    def __init__(self, source):
        super().__init__(convert_charrefs=True)
        self.source, self.stack, self.texts, self.blocks = source, [], [], []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        skip = tag in self.SKIP or a.get("translate") == "no" or "data-i18n-html" in a
        if "data-i18n-html" in a:
            inner = re.search(rf"{re.escape(self.get_starttag_text())}(.*?)</{tag}>", self.source, re.S).group(1)
            self.blocks.append(normal(inner))
        if tag not in ("input", "img", "meta", "link", "br"):
            self.stack.append(skip)
        if not any(self.stack):
            self.texts += [v for k, v in attrs if k in ATTRIBUTES and v and SPANISH.search(v)]

    def handle_endtag(self, tag):
        if tag not in ("input", "img", "meta", "link", "br") and self.stack:
            self.stack.pop()

    def handle_data(self, data):
        if not any(self.stack) and SPANISH.search(data):
            self.texts.append(normal(data))


@pytest.mark.parametrize("page", ["index.html", "mapeador.html", "modelos.html"])
def test_static_pages_are_translated(page):
    source = (ROOT / "web" / page).read_text(encoding="utf-8")
    p = _Page(source)
    p.feed(source)
    missing = [t for t in p.texts if translate(t) == t and t not in ("Multi Bridge", "MULTI", "BRIDGE")]
    assert not missing, f"{page}: add these to web/i18n/en.json: {missing}"
    assert not [b for b in p.blocks if b not in EN["html"]], f"{page}: a data-i18n-html block changed"


SKIP_KEYS = {"records", "samples", "content_b64", "example", "values", "before", "after", "upload_id", "id",
             "hash", "at", "key", "file", "name", "column", "path", "leaf", "concept", "format", "unit", "dtype",
             "label", "ai_view", "from", "fields", "proposal", "recipe", "approval", "source", "expected", "enum"}


WORDS = {"status", "proposed_by", "kind", "entity", "type", "op", "group", "role", "model", "pack", "group_by",
         "mime", "extension", "label", "format", "dtype", "unit", "from_unit", "to", "part"}   # codes, not sentences


def _file_text(path: Path) -> str:
    """Everything written in a file, to tell what is the file's own text (data) from the app's messages."""
    data = path.read_bytes()
    if path.suffix == ".xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(data))
        return " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None) \
            + " " + " ".join(c.comment.text for ws in wb for row in ws.iter_rows() for c in row if c.comment)
    for enc in ("utf-8", "cp1252"):
        try:
            return normal(data.decode(enc))
        except UnicodeDecodeError:
            pass
    return ""


def _messages(v, out, key=None):
    if isinstance(v, dict):
        for k, x in v.items():
            if k not in SKIP_KEYS:
                _messages(x, out, k)
    elif isinstance(v, list):
        for x in v:
            _messages(x, out, key)
    elif isinstance(v, str) and SPANISH.search(v) and key not in WORDS:
        out.add(v)


def test_what_the_api_says_is_translated(tmp_path):
    for folder in ("expected", "examples"):
        shutil.copytree(ROOT / folder, tmp_path / folder)
    app = App(tmp_path)
    dests = [d["key"] for d in route(app, "GET", "/api/destinations")]
    missing = set()
    for e in route(app, "GET", "/api/examples"):
        said: set[str] = set()
        up = route(app, "POST", "/api/upload", {"example": e["name"]})
        _messages(up, said)
        for d in dests:
            try:
                p = route(app, "POST", "/api/propose", {"upload_id": up["upload_id"], "destination": d})
            except ApiError as ex:
                said.add(str(ex))
                continue
            _messages({k: v for k, v in p.items() if k != "draft"}, said)
            _messages(p["draft"]["fields"], said)
            decisions = {k: {"status": f["status"], "from": f["from"], "transforms": f["transforms"]}
                         for k, f in p["draft"]["fields"].items()}
            try:
                route(app, "POST", "/api/approve", {"upload_id": up["upload_id"], "destination": d,
                                                    "draft": {"fields": decisions}, "by": ""})
            except ApiError as ex:
                said.add(str(ex))
        own = _file_text(tmp_path / "examples" / e["name"])
        for s in said:   # the file's own words (titles, comments, headers) are data: never translated
            if translate(s) == normal(s) and normal(s) not in own and normal(s.split(": ", 1)[-1]) not in own:
                missing.add(s)
    assert not missing, f"add these to web/i18n/en.json: {sorted(missing)}"


def test_data_patterns_dont_touch_names():
    assert translate("Hospital Universitario de León") == "Hospital Universitario de León"
    assert translate("3 de 3 hospitales de ejemplo conectados · 1 receta(s)") == "3 of 3 example hospitals connected · 1 recipe(s)"
    assert translate("Identificador del paciente: ID — solo donde PID-3.5 = MR (número de historia clínica)") \
        == "Patient identifier: ID — only where PID-3.5 = MR (medical record number)"
